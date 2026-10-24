import copy
import json

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def call(name, function, *args):
    print(json.dumps({'type': 'script.call', 'name': name}), flush=True)
    return function(*args)


def format_cell(cell):
    font = copy.copy(cell.font)
    font.bold = True
    cell.font = font
    cell.fill = PatternFill('solid', fgColor='FFD9EAF7')


workbook = call('openpyxl.load_workbook', load_workbook, 'input.xlsx')
for cell in workbook['Sales'][1][:3]:
    call('cell.format', format_cell, cell)
call('workbook.save', workbook.save, 'output.xlsx')
