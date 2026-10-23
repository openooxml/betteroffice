import importlib.util
import hashlib
import json
import os
import re
import shlex
import subprocess
import sys
import tempfile
import time
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

ROOT = Path(__file__).resolve().parent


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


runner = load('format_runner', 'run.py')
aggregator = load('format_aggregate', 'aggregate.py')
validation = load('format_validation', 'validate-checkers.py')
FAKE_CODEX = '''#!/usr/bin/env python3
import json
import os
import shutil
import sys
import time
from pathlib import Path
args = sys.argv[1:]
assert args[:2] == ['exec', '-m']
assert args[args.index('--sandbox') + 1] == 'read-only'
assert '--skip-git-repo-check' in args and '--ignore-user-config' in args
assert '--ignore-rules' in args and '--ephemeral' in args and '--json' in args
assert sys.stdin.read() == ''
workspace = Path(args[args.index('-C') + 1])
assert Path.cwd() == workspace
assert len(list(workspace.iterdir())) == 1
assert workspace.name.startswith('office-format-')
if os.environ.get('FAKE_WORKSPACES'):
    with open(os.environ['FAKE_WORKSPACES'], 'a') as stream:
        stream.write(json.dumps({'workspace': str(workspace), 'start': time.monotonic()}) + '\\n')
    deadline = time.monotonic() + 5
    while os.environ.get('FAKE_BARRIER') and time.monotonic() < deadline:
        markers = [json.loads(line) for line in Path(os.environ['FAKE_WORKSPACES']).read_text().splitlines()]
        if sum('start' in marker for marker in markers) >= 2:
            break
        time.sleep(0.01)
    time.sleep(0.2)
    with open(os.environ['FAKE_WORKSPACES'], 'a') as stream:
        stream.write(json.dumps({'workspace': str(workspace), 'end': time.monotonic()}) + '\\n')
overrides = [args[i + 1] for i, arg in enumerate(args) if arg == '-c']
assert 'approval_policy="never"' in overrides
assert 'mcp_servers.betteroffice.required=true' in overrides
assert 'mcp_servers.betteroffice.startup_timeout_sec=' + os.environ['FAKE_STARTUP_TIMEOUT'] in overrides
command = json.loads(next(value.split('=', 1)[1] for value in overrides if value.startswith('mcp_servers.betteroffice.command=')))
mcp_args = json.loads(next(value.split('=', 1)[1] for value in overrides if value.startswith('mcp_servers.betteroffice.args=')))
assert Path(command).is_absolute() and Path(command).is_file()
assert mcp_args[-2:] == ['--root', str(workspace)]
assert mcp_args[:-2] == ['argument with spaces']
assert args[-1].endswith('Use only the betteroffice MCP tools to read and edit the documents. All other tool calls disqualify the task. The documents live in the current directory.')
tasks = json.loads(Path(os.environ['FAKE_TASKS']).read_text())
task = next(task for task in tasks if args[-1].startswith(task['prompt']))
assert args[2] == os.environ['FAKE_MODEL']
if os.environ.get('FAKE_MODE') == 'hang':
    time.sleep(60)
if os.environ.get('FAKE_MODE') == 'crash':
    sys.exit(9)
if os.environ.get('FAKE_MODE') != 'missing':
    source = workspace / task['input'] if os.environ.get('FAKE_MODE') == 'untouched' else Path(os.environ['FAKE_CORRECT']) / (task['id'] + '.' + task['format'])
    shutil.copyfile(source, workspace / task['output'])
if os.environ.get('FAKE_MODE') == 'change-source':
    (workspace / task['input']).write_bytes(b'changed')
mode = os.environ.get('FAKE_MODE')
if mode == 'empty-log':
    sys.exit(0)
if mode == 'unstructured-log':
    print('Done')
    sys.exit(0)
if mode == 'message-only':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'message-1', 'type': 'agent_message', 'text': 'Done'}}))
    sys.exit(0)
print(json.dumps({'type': 'thread.started'}))
mode = os.environ.get('FAKE_MODE')
for event in [] if mode == 'no-tools' else ['item.started', 'item.completed']:
    print(json.dumps({'type': event, 'item': {'id': 'tool-1', 'type': 'mcp_tool_call', 'server': 'betteroffice'}}))
if mode == 'command':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'shell-1', 'type': 'command_execution', 'command': 'unzip -p input.xlsx'}}))
elif mode == 'other-mcp':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'other-tool', 'type': 'mcp_tool_call', 'server': 'other'}}))
elif mode == 'other-tool':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'browser-1', 'type': 'web_search', 'query': 'file:///input.xlsx'}}))
'''
SERVER_TOOLS = ['office_files', 'office_open', 'office_outline', 'office_grep', 'office_read', 'office_cells', 'office_render',
                'office_verify', 'office_close', 'office_propose', 'office_propose_cells', 'office_review', 'office_accept',
                'office_reject', 'office_export', 'xlsx_outline', 'xlsx_read_range', 'xlsx_preview', 'xlsx_write_range',
                'xlsx_edit_grid', 'xlsx_edit_sheets', 'xlsx_find_replace', 'xlsx_format', 'xlsx_sort',
                'pptx_outline', 'pptx_read_slide', 'pptx_preview', 'pptx_edit', 'docx_edit']
FAKE_MCP = '''#!/usr/bin/env python3
import json
import os
import sys
from pathlib import Path
args = sys.argv[1:]
assert args[:-2] == ['argument with spaces'] and args[-2] == '--root'
assert Path(args[-1]).is_dir() and not any(Path(args[-1]).iterdir())
mode = os.environ.get('FAKE_MCP_MODE')
if mode == 'exit':
    sys.exit(3)
tools = [name for name in json.loads(os.environ['FAKE_MCP_TOOLS']) if not (mode == 'missing-tool' and name == 'xlsx_sort')]
initialized = False
for line in sys.stdin:
    message = json.loads(line)
    if message['method'] == 'notifications/initialized':
        initialized = True
        continue
    if message['method'] == 'initialize':
        result = {'protocolVersion': message['params']['protocolVersion'], 'capabilities': {'tools': {}},
                  'serverInfo': {'name': 'fake', 'version': '1'}}
    else:
        assert message['method'] == 'tools/list' and initialized
        if mode == 'silent':
            continue
        start = int(message['params'].get('cursor', '0'))
        result = {'tools': [{'name': name, 'inputSchema': {'type': 'object'}} for name in tools[start:start + 10]]}
        if start + 10 < len(tools):
            result['nextCursor'] = str(start + 10)
    print(json.dumps({'jsonrpc': '2.0', 'id': message['id'], 'result': result}), flush=True)
'''
GENERIC_SUFFIX = (' Read and edit the documents however you like: by hand (for example unzip the package, edit its XML parts and zip it again), with any shell tool, or with Python; openpyxl, python-pptx and python-docx are installed but optional.'
                  ' There is no network access. All other tool calls disqualify the task. The documents live in the current directory.'
                  ' Work only inside this directory; temporary files may go in it.'
                  ' Commands that reference the repository, evaluation sources or task metadata, and file changes outside this directory disqualify the task.')
FAKE_GENERIC_CODEX = '''#!/usr/bin/env python3
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
args = sys.argv[1:]
assert args[:3] == ['exec', '-m', os.environ['FAKE_MODEL']]
assert args[args.index('--sandbox') + 1] == 'workspace-write'
workspace = Path(args[args.index('-C') + 1])
assert Path.cwd() == workspace and len(list(workspace.iterdir())) == 1
assert not workspace.is_relative_to(Path(os.environ['FAKE_REPO']))
assert os.environ['PATH'].split(os.pathsep)[0] == os.environ['FAKE_PYTHON_BIN']
overrides = [args[i + 1] for i, arg in enumerate(args) if arg == '-c']
assert 'sandbox_workspace_write.network_access=false' in overrides
assert not any(value.startswith(('mcp_servers', 'sandbox_workspace_write.writable_roots')) for value in overrides)
assert args[-1].endswith(os.environ['FAKE_GENERIC_SUFFIX'])
task = next(task for task in json.loads(Path(os.environ['FAKE_TASKS']).read_text()) if args[-1].startswith(task['prompt']))
mode = os.environ.get('FAKE_MODE')
if mode == 'hand-edit':
    scratch = workspace / 'work'
    with ZipFile(workspace / task['input']) as archive:
        archive.extractall(scratch)
    subprocess.run("sed 's/Sumary/Summary/' work/word/document.xml > work/document.xml && mv work/document.xml work/word/document.xml",
                   shell=True, check=True)
    with ZipFile(workspace / task['output'], 'w', compression=ZIP_DEFLATED) as archive:
        for path in sorted(scratch.rglob('*')):
            if path.is_file():
                archive.write(path, path.relative_to(scratch).as_posix())
else:
    shutil.copyfile(Path(os.environ['FAKE_CORRECT']) / (task['id'] + '.' + task['format']), workspace / task['output'])
if mode == 'empty-log':
    sys.exit(0)
if mode == 'unstructured-log':
    print('Done')
    sys.exit(0)
if mode == 'message-only':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'message-1', 'type': 'agent_message', 'text': 'Done'}}))
    sys.exit(0)
print(json.dumps({'type': 'thread.started'}))
items = [
    {'id': 'shell-1', 'type': 'command_execution', 'command': 'unzip -o input.docx -d work/'},
    {'id': 'shell-2', 'type': 'command_execution',
     'command': "sed 's/Sumary/Summary/' work/word/document.xml > work/document.xml && mv work/document.xml work/word/document.xml"},
    {'id': 'shell-3', 'type': 'command_execution', 'command': 'cd work && zip -X -r ../output.docx .'},
] if mode == 'hand-edit' else [] if mode == 'no-tools' else [
    {'id': 'shell-1', 'type': 'command_execution', 'command': 'python3 edit.py'},
    {'id': 'patch-1', 'type': 'file_change', 'changes': [{'path': 'edit.py', 'kind': 'add'}]}]
for item in items:
    for event in ['item.started', 'item.completed']:
        print(json.dumps({'type': event, 'item': item}))
mode = os.environ.get('FAKE_MODE')
if mode == 'mcp':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'tool-1', 'type': 'mcp_tool_call', 'server': 'betteroffice'}}))
elif mode == 'web':
    print(json.dumps({'type': 'item.completed', 'item': {'id': 'web-1', 'type': 'web_search', 'query': 'openpyxl docs'}}))
elif mode == 'path':
    command = os.environ.get('FAKE_COMMAND', '').replace('__REPOSITORY_FROM_WORKSPACE__', os.path.relpath(os.environ['FAKE_REPO'], workspace))
    item = {'id': 'outside-1', 'type': 'command_execution', 'command': command} if 'FAKE_COMMAND' in os.environ else {
        'id': 'outside-1', 'type': 'file_change', 'changes': [json.loads(os.environ['FAKE_CHANGE'])]}
    print(json.dumps({'type': 'item.completed', 'item': item}))
elif mode == 'task-file':
    name = os.environ['FAKE_TASK_FILE']
    (workspace / name).write_text('pass')
    if os.environ['FAKE_REPORT_CHANGE'] == '1':
        print(json.dumps({'type': 'item.completed', 'item': {
            'id': 'task-file-1', 'type': 'file_change', 'changes': [{'path': name, 'kind': 'add'}]}}))
    print(json.dumps({'type': 'item.completed', 'item': {
        'id': 'local-1', 'type': 'command_execution', 'command': os.environ['FAKE_COMMAND']}}))
'''

SCRIPT_MCP = '''#!/usr/bin/env python3
import json
import os
import shutil
import sys
from pathlib import Path
root = Path(sys.argv[-1])
initialized = False
for line in sys.stdin:
    message = json.loads(line)
    method = message['method']
    if method == 'notifications/initialized':
        initialized = True
        continue
    if method == 'initialize':
        result = {'protocolVersion': message['params']['protocolVersion'], 'capabilities': {'tools': {}},
                  'serverInfo': {'name': 'script-stub', 'version': '1'}}
    elif method == 'tools/list':
        assert initialized
        result = {'tools': [{'name': name} for name in json.loads(os.environ['FAKE_MCP_TOOLS'])]}
    else:
        assert method == 'tools/call' and initialized
        params = message['params']
        name, args = params['name'], params['arguments']
        with open(os.environ['SCRIPT_CALLS'], 'a') as stream:
            stream.write(json.dumps({'root': str(root), 'name': name, 'arguments': args}) + '\\n')
        assert len(list(root.glob('input.*'))) == 1
        if name == 'office_open':
            result = {'document': 'random-document-id', 'version': 'random-version'}
        else:
            assert args['document'] == 'random-document-id'
            if name == 'office_grep':
                assert args['query'] == 'Sumary' and args['caseSensitive'] is True
                result = {'matches': [{'match': 'random-match-id'}]}
            elif name == 'office_propose':
                assert args['edits'] == [{'match': 'random-match-id', 'newText': 'Summary'}]
                result = {'id': 'random-proposal-id'}
            elif name == 'xlsx_format':
                assert args['version'] == 'random-version'
                assert args['edits'] == [{'sheet': 'Sales', 'range': 'A1:C1', 'bold': True, 'fill': '#D9EAF7'}]
                result = {'version': 'after-format'}
            elif name == 'office_export':
                ext = Path(args['path']).suffix
                task = 'docx-01-fix-heading' if ext == '.docx' else 'xlsx-04-style-header'
                if ext == '.docx':
                    assert args['proposal'] == 'random-proposal-id'
                shutil.copyfile(Path(os.environ['FAKE_CORRECT']) / (task + ext), root / args['path'])
                result = {'path': args['path']}
            else:
                raise AssertionError(name)
        result = {'structuredContent': result, 'content': []}
        if os.environ.get('SCRIPT_FAIL') == name:
            result = {'isError': True, 'content': [{'type': 'text', 'text': 'refused'}]}
    print(json.dumps({'jsonrpc': '2.0', 'id': message['id'], 'result': result}), flush=True)
'''


def executable(path, text):
    path.parent.mkdir(exist_ok=True)
    path.write_text(text)
    path.chmod(0o755)
    return path


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='office-runner-test-')
        self.root = Path(self.temporary.name)
        self.bin = self.root / 'bin with spaces'
        self.bin.mkdir()
        fake = self.bin / 'codex'
        fake.write_text(FAKE_CODEX)
        fake.chmod(0o755)
        cli = self.bin / 'mcp'
        cli.write_text(FAKE_MCP)
        cli.chmod(0o755)
        self.cli = [str(cli), 'argument with spaces']
        self.correct = self.root / 'correct'
        self.correct.mkdir()
        for task in runner.checker.tasks():
            package = validation.Package(runner.checker.fixture(task))
            number = int(task['id'].split('-')[1])
            validation.CORRECT[task['format']](package, number)
            package.save(self.correct / f'{task["id"]}.{task["format"]}')
        self.environment = patch.dict(os.environ, {
            'PATH': str(self.bin) + os.pathsep + os.environ['PATH'], 'FAKE_TASKS': str(ROOT / 'tasks.json'),
            'FAKE_CORRECT': str(self.correct), 'FAKE_MODEL': 'test-model', 'FAKE_STARTUP_TIMEOUT': '60',
            'FAKE_MCP_TOOLS': json.dumps(SERVER_TOOLS),
        })
        self.environment.start()
        self.generic_bin = self.root / 'generic bin'
        executable(self.generic_bin / 'codex', FAKE_GENERIC_CODEX)
        self.python = executable(self.root / 'python bin' / 'python',
                                 '#!/bin/sh\n[ "$1" = -c ] && [ "$2" = "import openpyxl, pptx, docx" ] || exit 7\n')
        self.broken_python = executable(self.root / 'broken python' / 'python',
                                        "#!/bin/sh\necho \"ModuleNotFoundError: No module named 'docx'\" >&2\nexit 1\n")

    def generic_environment(self, **values):
        return patch.dict(os.environ, {'PATH': str(self.generic_bin) + os.pathsep + os.environ['PATH'],
                                       'FAKE_PYTHON_BIN': str(self.python.parent), 'FAKE_GENERIC_SUFFIX': GENERIC_SUFFIX,
                                       'FAKE_REPO': str(runner.REPO), **values})

    def tearDown(self):
        self.environment.stop()
        self.temporary.cleanup()

    def run_one(self, mode):
        results = self.root / mode
        results.mkdir()
        with patch.dict(os.environ, {'FAKE_MODE': mode}):
            return runner.run_task(runner.checker.tasks()[0], self.cli, 'test-model', results)

    def test_full_shell_run_and_reports(self):
        results = self.root / 'results with spaces'
        command = ' '.join(shlex.quote(value) for value in self.cli)
        result = subprocess.run([str(ROOT / 'run.sh'), command, str(results), 'test-model'],
                                stdin=subprocess.DEVNULL, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual(len(report['results']), 30)
        self.assertEqual(report['timeout_seconds'], 600)
        self.assertGreater(report['mcp_startup_seconds'], 0)
        self.assertEqual(report['arm'], 'betteroffice')
        self.assertEqual(report['formats'], ['xlsx', 'pptx', 'docx'])
        self.assertTrue(all(item['pass'] and item['tool_calls'] == 1 and item['duration_seconds'] > 0 for item in report['results']))
        self.assertEqual((results / 'results.md').read_text().count('| PASS |'), 30)
        self.assertIn('Formats: xlsx, pptx, docx.', (results / 'results.md').read_text())
        repeat = subprocess.run([str(ROOT / 'run.sh'), command, str(results), 'test-model'], capture_output=True)
        self.assertEqual(repeat.returncode, 2)

    def test_formats_select_the_tasks_and_the_preflight_tools(self):
        command = ' '.join(shlex.quote(value) for value in self.cli)
        with patch.dict(os.environ, {'FAKE_MCP_MODE': 'missing-tool'}):
            selected = subprocess.run([str(ROOT / 'run.sh'), '--formats', 'docx,pptx', command, str(self.root / 'selected'), 'test-model'],
                                      stdin=subprocess.DEVNULL, capture_output=True, text=True)
            xlsx = subprocess.run([str(ROOT / 'run.sh'), '--formats', 'xlsx', command, str(self.root / 'xlsx'), 'test-model'],
                                  stdin=subprocess.DEVNULL, capture_output=True, text=True)
        self.assertEqual(selected.returncode, 0, selected.stdout + selected.stderr)
        report = json.loads((self.root / 'selected' / 'results.json').read_text())
        expected = [task['id'] for task in runner.checker.tasks() if task['format'] in ('pptx', 'docx')]
        self.assertEqual(len(expected), 20)
        self.assertEqual([item['task'] for item in report['results']], expected)
        self.assertEqual(report['formats'], ['pptx', 'docx'])
        self.assertTrue(all(item['pass'] for item in report['results']))
        self.assertIn('Formats: pptx, docx. Passed: 20/20.', (self.root / 'selected' / 'results.md').read_text())
        self.assertEqual(xlsx.returncode, 1, xlsx.stdout + xlsx.stderr)
        self.assertIn('MCP server not ready: missing required tools: xlsx_sort', xlsx.stderr)

    def test_mcp_preflight_requires_a_ready_server_with_the_task_tools(self):
        required = runner.required_tools(runner.checker.tasks())
        self.assertTrue({'office_open', 'office_export', 'xlsx_sort', 'pptx_edit', 'docx_edit'} <= set(required) <= set(SERVER_TOOLS))
        self.assertGreater(runner.mcp_ready(self.cli, required), 0)
        for mode, reason in [('silent', 'did not answer tools/list within 1 seconds'),
                             ('missing-tool', 'missing required tools: xlsx_sort'), ('exit', 'exited with code 3')]:
            with self.subTest(mode=mode), patch.dict(os.environ, {'FAKE_MCP_MODE': mode}), \
                    patch.object(runner, 'STARTUP_TIMEOUT_SECONDS', 1):
                started = time.monotonic()
                with self.assertRaisesRegex(RuntimeError, 'MCP server not ready: .*' + re.escape(reason)):
                    runner.mcp_ready(self.cli, required)
                self.assertLess(time.monotonic() - started, 5)

    def test_run_stops_before_the_first_task_when_the_server_is_not_ready(self):
        results = self.root / 'not ready'
        command = ' '.join(shlex.quote(value) for value in self.cli)
        with patch.dict(os.environ, {'FAKE_MCP_MODE': 'missing-tool'}):
            result = subprocess.run([str(ROOT / 'run.sh'), command, str(results), 'test-model'],
                                    stdin=subprocess.DEVNULL, capture_output=True, text=True)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        self.assertIn('MCP server not ready: missing required tools: xlsx_sort', result.stderr)
        self.assertFalse(results.exists())

    def test_codex_waits_for_the_required_server(self):
        command = runner.command_line(self.cli, self.root, 'test-model', 'Do the task.')
        overrides = [command[i + 1] for i, arg in enumerate(command) if arg == '-c']
        self.assertIn('mcp_servers.betteroffice.required=true', overrides)
        self.assertIn('mcp_servers.betteroffice.startup_timeout_sec=60', overrides)

    def test_tasks_without_tool_calls_are_reported(self):
        record = self.run_one('no-tools')
        runner.checker.check(runner.checker.tasks()[0], self.root / 'no-tools' / record['task'] / 'output.xlsx')
        self.assertEqual(record['tool_calls'], 0)
        self.assertFalse(record['pass'])
        self.assertIn('no tool calls', record['check_message'])
        runner.write_results(self.root / 'no-tools', 'test-model', [record, self.run_one('missing')])
        report = (self.root / 'no-tools' / 'results.md').read_text()
        self.assertIn('Tasks without tool calls: 1.', report)
        self.assertRegex(report, r'\| xlsx-01-correct-value \| FAIL \| [^|]+ \| 0 \|')

    def test_generic_tasks_without_tool_calls_fail_with_correct_output(self):
        results = self.root / 'generic-no-tools'
        results.mkdir()
        with self.generic_environment(FAKE_MODE='no-tools'):
            record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
        runner.checker.check(runner.checker.tasks()[0], results / record['task'] / 'output.xlsx')
        self.assertEqual(record['tool_calls'], 0)
        self.assertFalse(record['pass'])
        self.assertIn('no tool calls', record['check_message'])

    def test_unavailable_tool_counts_fail_with_correct_output_in_both_arms(self):
        task = runner.checker.tasks()[0]
        for arm in ('betteroffice', 'generic'):
            for mode in ('empty-log', 'unstructured-log', 'message-only'):
                results = self.root / f'{arm}-{mode}'
                results.mkdir()
                environment = self.generic_environment(FAKE_MODE=mode) if arm == 'generic' else patch.dict(os.environ, {'FAKE_MODE': mode})
                with self.subTest(arm=arm, mode=mode), environment:
                    record = runner.run_task(task, self.cli if arm == 'betteroffice' else None, 'test-model', results,
                                             arm=arm, python=self.python if arm == 'generic' else None)
                    runner.checker.check(task, results / record['task'] / task['output'])
                    self.assertEqual(record['agent_exit_code'], 0)
                    self.assertIsNone(record['tool_calls'])
                    self.assertFalse(record['pass'])
                    self.assertEqual(record['check_message'], 'agent made no tool calls')

    def test_missing_output_crash_and_untouched_export_fail(self):
        for mode in ['missing', 'crash', 'untouched', 'change-source']:
            record = self.run_one(mode)
            self.assertFalse(record['pass'], mode)
            if mode == 'crash':
                self.assertEqual(record['agent_exit_code'], 9)
                self.assertIn('agent exited with code 9', record['check_message'])
            if mode == 'change-source':
                self.assertIn('input fixture was changed', record['check_message'])

    def test_timeout(self):
        with patch.object(runner, 'TIMEOUT_SECONDS', 0.1):
            record = self.run_one('hang')
        self.assertTrue(record['timed_out'])
        self.assertFalse(record['pass'])
        self.assertIn('timed out', record['check_message'])
        self.assertLess(record['duration_seconds'], 5)

    def test_forbidden_tools_fail_with_correct_output(self):
        for mode, reason in [('command', 'command_execution'), ('other-mcp', 'MCP server other'),
                             ('other-tool', 'web_search')]:
            with self.subTest(mode=mode):
                record = self.run_one(mode)
                self.assertEqual(record['agent_exit_code'], 0)
                runner.checker.check(runner.checker.tasks()[0], self.root / mode / record['task'] / 'output.xlsx')
                self.assertFalse(record['pass'])
                self.assertIn(reason, record['check_message'])
                self.assertEqual(record['tool_calls'], 1)
                runner.write_results(self.root / mode, 'test-model', [record])
                self.assertIn(reason, (self.root / mode / 'results.json').read_text())
                self.assertIn(reason, (self.root / mode / 'results.md').read_text())

    def test_file_access_features_disabled(self):
        command = runner.command_line(self.cli, self.root, 'test-model', 'Do the task.')
        disabled = {command[i + 1] for i, arg in enumerate(command) if arg == '--disable'}
        self.assertTrue({'shell_tool', 'browser_use', 'browser_use_external', 'browser_use_full_cdp_access',
                         'in_app_browser', 'computer_use', 'apps', 'view_image', 'in_app_local_automation',
                         'plugins', 'remote_plugin', 'multi_agent', 'hooks',
                         'skill_search', 'skill_mcp_dependency_install', 'image_generation',
                         'unified_exec_tty', 'shell_snapshot', 'workspace_dependencies',
                         'worktrees'} <= disabled)
        self.assertNotIn('code_mode_host', disabled)
        self.assertIn('--ignore-user-config', command)
        overrides = [command[i + 1] for i, arg in enumerate(command) if arg == '-c']
        self.assertIn('web_search="disabled"', overrides)

    def test_default_model_and_unavailable_tool_count(self):
        command = runner.command_line(self.cli, self.root, 'gpt-6-luna', 'Do the task.')
        self.assertEqual(command[command.index('-m') + 1], 'gpt-6-luna')
        log = self.root / 'plain.log'
        log.write_text('Not a structured Codex event\n[]\nnull\n')
        self.assertIsNone(runner.tool_calls(log))

    def test_generic_arm_command_allows_only_the_sandboxed_shell(self):
        command = runner.command_line(None, self.root, 'test-model', 'Do the task.', arm='generic')
        self.assertEqual(command[command.index('--sandbox') + 1], 'workspace-write')
        self.assertEqual(command[command.index('-C') + 1], str(self.root))
        self.assertFalse({'danger-full-access', '--add-dir', '--dangerously-bypass-approvals-and-sandbox'} & set(command))
        self.assertTrue({'--skip-git-repo-check', '--ignore-user-config', '--ignore-rules', '--ephemeral', '--json'} <= set(command))
        overrides = [command[i + 1] for i, arg in enumerate(command) if arg == '-c']
        self.assertEqual(sorted(overrides), ['approval_policy="never"', 'sandbox_workspace_write.network_access=false',
                                             'web_search="disabled"'])
        enabled = {command[i + 1] for i, arg in enumerate(command) if arg == '--enable'}
        disabled = {command[i + 1] for i, arg in enumerate(command) if arg == '--disable'}
        self.assertEqual(enabled, {'shell_tool', 'unified_exec', 'unified_exec_tty'})
        self.assertEqual(disabled, {'shell_snapshot', 'browser_use', 'browser_use_external', 'browser_use_full_cdp_access',
                                    'in_app_browser', 'computer_use', 'apps', 'view_image', 'in_app_local_automation',
                                    'plugins', 'remote_plugin', 'multi_agent', 'hooks', 'skill_search',
                                    'skill_mcp_dependency_install', 'image_generation', 'workspace_dependencies', 'worktrees'})
        self.assertEqual(command[-1], 'Do the task.' + GENERIC_SUFFIX)

    def test_generic_arm_counts_shell_work_and_fails_mcp_and_web_search(self):
        for mode, reason in [('', None), ('mcp', 'forbidden tool call: MCP server betteroffice'),
                             ('web', 'forbidden tool call: web_search')]:
            results = self.root / f'generic-{mode or "shell"}'
            results.mkdir()
            with self.subTest(mode=mode), self.generic_environment(FAKE_MODE=mode):
                record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
                self.assertEqual(record['tool_calls'], 2)
                self.assertEqual(record['pass'], reason is None, record['check_message'])
                if reason:
                    self.assertIn(reason, record['check_message'])

    def test_generic_hand_edit_in_scratch_directory_passes(self):
        results = self.root / 'generic-hand-edit'
        results.mkdir()
        task = next(task for task in runner.checker.tasks() if task['id'] == 'docx-01-fix-heading')
        with self.generic_environment(FAKE_MODE='hand-edit'):
            record = runner.run_task(task, None, 'test-model', results, arm='generic', python=self.python)
        runner.checker.check(task, results / task['id'] / task['output'])
        self.assertEqual(record['agent_exit_code'], 0)
        self.assertTrue(record['pass'], record['check_message'])
        self.assertEqual(record['tool_calls'], 3)

    def test_generic_repository_reads_and_outside_file_changes_fail_with_correct_output(self):
        home_relative = os.path.relpath(runner.REPO, Path.home())
        commands = [
            'cat eval/formats/check.py',
            'python3 ../other/eval/formats/check.py',
            f'cat {shlex.quote(str(ROOT / "check.py"))}',
            f'ls {shlex.quote(str(runner.REPO))}',
            'python3 -c "import eval.formats.check"',
            'python3 -c "from eval.formats import check"',
            'cat validate-checkers.py tasks.json fixtures.json',
            'cat tasks.json',
            f'cat {shlex.quote(str(runner.REPO / "AGENTS.md"))}',
            f'cat ~/{shlex.quote(home_relative + "/eval/formats/tasks.json")}',
            f'cat ~/{shlex.quote(home_relative + "/AGENTS.md")}',
            f'cat "$HOME/{home_relative}/AGENTS.md"',
            f'cat "${{HOME}}/{home_relative}/AGENTS.md"',
            'cat "__REPOSITORY_FROM_WORKSPACE__/AGENTS.md"',
        ]
        changes = [
            {'path': '../edit.py', 'kind': 'add'},
            {'path': '/outside/edit.py', 'kind': 'add'},
            {'path': str(runner.REPO / 'edit.py'), 'kind': 'add'},
            {'path': 'edit.py', 'kind': 'update', 'diff': '+ from eval.formats import check'},
        ]
        cases = [{'FAKE_COMMAND': command} for command in commands] + [{'FAKE_CHANGE': json.dumps(change)} for change in changes]
        for index, values in enumerate(cases):
            results = self.root / f'generic-path-{index}'
            results.mkdir()
            with self.subTest(values=values), self.generic_environment(FAKE_MODE='path', **values):
                record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
                runner.checker.check(runner.checker.tasks()[0], results / record['task'] / 'output.xlsx')
                self.assertEqual(record['agent_exit_code'], 0)
                self.assertFalse(record['pass'])
                self.assertIn('forbidden path access', record['check_message'])

    def test_generic_task_code_and_unrelated_paths_pass_with_correct_output(self):
        commands = [
            "python3 - <<'PY'\nwb.save('eval.xlsx')\nPY",
            "python3 - <<'PY'\nx = eval('1+1')\nPY",
            "python3 - <<'PY'\nxml = b'<a:t>Draft results</a:t>'\nPY",
            "python3 - <<'PY'\nxml = b'<c r=\"B5\" s=\"1\">...</c>'\nPY",
            "python3 - <<'PY'\nformula = f'=(C{r}-B{r})/B{r}'\nPY",
            "python3 - <<'PY'\nformula = '=(C2-B2)/B2'\nPY",
            'libreoffice --headless --convert-to xlsx input.xlsx >/tmp/office-recalc.log 2>&1',
            'python3 edit.py 2>/dev/null',
            "python3 - <<'PY'\nparts = name.split('/')\nPY",
            "python3 - <<'PY'\nx = ~mask\nPY",
            "python3 - <<'PY'\nformula = '=$B2/$C2'\nPY",
            'cat ../other/input.xlsx',
            'cat /etc/passwd',
            'cat "/outside directory/input.xlsx"',
            'cat ~/input.xlsx',
            'cat "$HOME/input.xlsx"',
        ]
        changes = [{'path': 'edit.py', 'kind': 'update', 'diff': '+ text = open("/etc/passwd").read()'}]
        cases = [{'FAKE_COMMAND': command} for command in commands] + [{'FAKE_CHANGE': json.dumps(change)} for change in changes]
        home = self.root / 'home'
        home.mkdir()
        for index, values in enumerate(cases):
            results = self.root / f'generic-task-code-{index}'
            results.mkdir()
            with self.subTest(values=values), self.generic_environment(FAKE_MODE='path', HOME=str(home), **values):
                record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
                runner.checker.check(runner.checker.tasks()[0], results / record['task'] / 'output.xlsx')
                self.assertEqual(record['agent_exit_code'], 0)
                self.assertTrue(record['pass'], record['check_message'])

    def test_generic_task_local_evaluation_names_pass_with_correct_output(self):
        names = ['check.py', 'validate-checkers.py', 'tasks.json', 'fixtures.json']
        for name in names:
            for report_change in ('0', '1'):
                results = self.root / f'generic-local-{name}-{report_change}'
                results.mkdir()
                command = f'python3 {name}' if name.endswith('.py') else f'cat {name}'
                with self.subTest(name=name, report_change=report_change), self.generic_environment(
                        FAKE_MODE='task-file', FAKE_TASK_FILE=name, FAKE_REPORT_CHANGE=report_change, FAKE_COMMAND=command):
                    record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
                    runner.checker.check(runner.checker.tasks()[0], results / record['task'] / 'output.xlsx')
                    self.assertEqual(record['agent_exit_code'], 0)
                    self.assertTrue(record['pass'], record['check_message'])
                    self.assertEqual(record['tool_calls'], 3 + int(report_change))

    def test_generic_local_file_names_do_not_allow_evaluation_paths(self):
        workspace = self.root / 'workspace'
        workspace.mkdir()
        for name in ['check.py', 'validate-checkers.py', 'tasks.json', 'fixtures.json']:
            (workspace / name).write_text('pass')
        for command in ['python3 ../other/eval/formats/check.py', 'cat eval/formats/tasks.json',
                        'cat eval/other.py', 'python3 -c "import eval"', 'python3 -c "from eval import formats"',
                        'python3 -m eval.formats.check',
                        "python3 -c \"print(open('/srv/' + 'clone/' + 'eval/' + 'formats/check.py').read())\""]:
            with self.subTest(command=command):
                self.assertIn('forbidden path access', runner.path_violation(command, workspace))

    def test_generic_path_audit_checks_repository_spellings_anywhere(self):
        workspace = self.root / 'workspace'
        workspace.mkdir()
        repository = self.root / 'repository-link'
        repository.symlink_to(runner.REPO, target_is_directory=True)
        log = self.root / 'repository-paths.jsonl'
        for spelling in (repository, repository.resolve()):
            with self.subTest(spelling=spelling), patch.object(runner, 'REPO', repository):
                log.write_text(json.dumps({'type': 'item.completed', 'item': {
                    'type': 'command_execution', 'command': f"printf 'path={spelling}/file'"}}) + '\n')
                self.assertIn('forbidden path access', runner.access_violation(log, 'generic', workspace))

    def test_generic_path_audit_matches_repository_path_boundaries(self):
        workspace = self.root / 'workspace'
        workspace.mkdir()
        with patch.object(runner, 'REPO', Path('/work')):
            for command in ['python3 -c \'open("/tmp/x/workspace/input.xlsx")\'',
                            'cat /workspace/input.xlsx', 'cat /work-backup/input.xlsx',
                            'cat /tmp/x/work/input.xlsx']:
                with self.subTest(command=command):
                    self.assertIsNone(runner.path_violation(command, workspace))
            for command in ['cat /work/eval/x', 'cat /work/file', 'cat "/work"',
                            'printf "path=/work/file"']:
                with self.subTest(command=command):
                    self.assertIn('forbidden path access', runner.path_violation(command, workspace))

    def test_generic_workspace_stays_outside_repository_when_tempdir_is_in_repository(self):
        results = self.root / 'generic-tempdir'
        results.mkdir()
        with self.generic_environment(), patch.object(runner.tempfile, 'tempdir', str(runner.REPO)):
            record = runner.run_task(runner.checker.tasks()[0], None, 'test-model', results, arm='generic', python=self.python)
        self.assertTrue(record['pass'], record['check_message'])

    def test_generic_path_audit_allows_task_files_and_non_repository_paths(self):
        workspace = self.root / 'workspace'
        workspace.mkdir()
        log = self.root / 'paths.jsonl'
        library = self.root / 'runtime' / 'lib'
        commands = [
            'python3 edit.py',
            '/bin/sh -c "python3 edit.py"',
            'python3 -c \'from pathlib import Path; p = Path.cwd() / "input.xlsx"\'',
            'python3 -c \'ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"\'',
            f'python3 -c \'open("{workspace}/input.xlsx")\'',
            f'{shlex.quote(str(self.python))} edit.py',
            f'python3 -c \'open("{library}/python3.14/pathlib.py")\'',
            f'python3 -c \'open("{library}/python3.14/site-packages/docx/__init__.py")\'',
        ]
        for command in commands:
            with self.subTest(command=command):
                log.write_text(json.dumps({'type': 'item.completed', 'item': {
                    'type': 'command_execution', 'command': command}}) + '\n')
                self.assertIsNone(runner.access_violation(log, 'generic', workspace))
        (workspace / 'linked').symlink_to(runner.REPO, target_is_directory=True)
        log.write_text(json.dumps({'type': 'item.completed', 'item': {
            'type': 'command_execution', 'command': 'cat linked/AGENTS.md'}}) + '\n')
        self.assertIn('forbidden path access', runner.access_violation(log, 'generic', workspace))

    def test_generic_shell_run_checks_its_python(self):
        results = self.root / 'generic results'
        with self.generic_environment():
            result = subprocess.run([str(ROOT / 'run.sh'), '--arm', 'generic', '--python', str(self.python), str(results), 'test-model'],
                                    stdin=subprocess.DEVNULL, capture_output=True, text=True)
            broken = subprocess.run([str(ROOT / 'run.sh'), '--arm', 'generic', '--python', str(self.broken_python),
                                     str(self.root / 'broken'), 'test-model'], stdin=subprocess.DEVNULL, capture_output=True, text=True)
            usage = subprocess.run([str(ROOT / 'run.sh'), '--arm', 'generic', str(self.root / 'usage')],
                                   stdin=subprocess.DEVNULL, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual((report['arm'], report['python']), ('generic', str(self.python)))
        self.assertNotIn('mcp_startup_seconds', report)
        self.assertTrue(all(item['pass'] and item['tool_calls'] == 2 for item in report['results']))
        self.assertIn('Arm: `generic`', (results / 'results.md').read_text())
        self.assertEqual(broken.returncode, 1)
        self.assertIn('cannot import openpyxl, pptx and docx', broken.stderr)
        self.assertFalse((self.root / 'broken').exists())
        self.assertEqual(usage.returncode, 2)
        self.assertIn('--python', usage.stderr)

    def test_generic_formats_and_unknown_format(self):
        generic = [str(ROOT / 'run.sh'), '--arm', 'generic', '--python', str(self.python), '--formats']
        with self.generic_environment():
            result = subprocess.run([*generic, 'docx', str(self.root / 'generic docx'), 'test-model'],
                                    stdin=subprocess.DEVNULL, capture_output=True, text=True)
            unknown = subprocess.run([*generic, 'docx,odt', str(self.root / 'unknown'), 'test-model'],
                                     stdin=subprocess.DEVNULL, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        report = json.loads((self.root / 'generic docx' / 'results.json').read_text())
        self.assertEqual(report['formats'], ['docx'])
        self.assertEqual([item['task'] for item in report['results']],
                         [task['id'] for task in runner.checker.tasks() if task['format'] == 'docx'])
        self.assertTrue(all(item['pass'] and item['tool_calls'] == 2 for item in report['results']))
        self.assertEqual(unknown.returncode, 2)
        self.assertIn("argument --formats: unknown format in 'docx,odt'", unknown.stderr)
        self.assertFalse((self.root / 'unknown').exists())

    def invoke(self, *arguments):
        with patch.object(sys, 'argv', ['run.py', *map(str, arguments)]):
            return runner.main()

    def command(self):
        return ' '.join(shlex.quote(value) for value in self.cli)

    def test_repeated_runs_layout_and_concurrency(self):
        results, workspaces = self.root / 'repeated', self.root / 'workspaces.jsonl'
        tasks = runner.checker.tasks()[:2]
        with patch.dict(os.environ, {'FAKE_WORKSPACES': str(workspaces), 'FAKE_BARRIER': '1'}):
            status = self.invoke('--runs', '3', '--jobs', '2', '--tasks', ','.join(task['id'] for task in tasks),
                                 self.command(), results, 'test-model')
        self.assertEqual(status, 0)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual((report['suite'], report['runs'], report['agent']), ('v2', 3, 'codex'))
        self.assertEqual([(record['task'], record['run']) for record in report['results']],
                         [(task['id'], index) for task in tasks for index in range(1, 4)])
        for record in report['results']:
            directory = results / record['task'] / f'run-{record["run"]}'
            self.assertTrue((directory / 'codex.jsonl').is_file())
            self.assertTrue((directory / 'output.xlsx').is_file())
            self.assertEqual((record['passed'], record['reason'], record['seconds']),
                             (record['pass'], record['check_message'], record['duration_seconds']))
        markers = [json.loads(line) for line in workspaces.read_text().splitlines()]
        self.assertEqual(len({marker['workspace'] for marker in markers}), 6)
        active = peak = 0
        for marker in sorted(markers, key=lambda item: item.get('start', item.get('end'))):
            active += 1 if 'start' in marker else -1
            peak = max(peak, active)
        self.assertEqual(peak, 2)
        self.assertTrue(all(not Path(marker['workspace']).exists() for marker in markers))

    def test_resume_skips_finished_pairs_and_retries_unrecorded_directory(self):
        results, workspaces = self.root / 'resume', self.root / 'resume-workspaces.jsonl'
        task = runner.checker.tasks()[0]
        arguments = ['--runs', '2', '--tasks', task['id'], self.command(), results, 'test-model']
        with patch.dict(os.environ, {'FAKE_WORKSPACES': str(workspaces)}):
            self.assertEqual(self.invoke(*arguments), 0)
            report = json.loads((results / 'results.json').read_text())
            log = results / task['id'] / 'run-1' / 'codex.jsonl'
            old_log = log.read_bytes()
            report['results'].pop()
            (results / 'results.json').write_text(json.dumps(report))
            self.assertEqual(self.invoke('--resume', *arguments), 0)
        self.assertEqual(len(workspaces.read_text().splitlines()), 6)
        self.assertEqual(log.read_bytes(), old_log)
        self.assertEqual(len(json.loads((results / 'results.json').read_text())['results']), 2)
        with patch.object(runner, 'run_task', side_effect=AssertionError('finished runs must be skipped')), \
                patch.object(runner, 'mcp_ready', side_effect=AssertionError('no pending runs')):
            self.assertEqual(self.invoke('--resume', *arguments), 0)
        with self.assertRaises(SystemExit) as error:
            self.invoke('--resume', '--runs', '3', '--tasks', task['id'], self.command(), results, 'test-model')
        self.assertEqual(error.exception.code, 2)

    def test_resume_keeps_finished_failures_and_saves_incrementally(self):
        results = self.root / 'resume-failure'
        tasks = runner.checker.tasks()[:2]
        original = runner.run_task

        def crash(task, *args, **kwargs):
            if task['id'] == tasks[1]['id']:
                deadline = time.monotonic() + 5
                saved = json.loads((results / 'results.json').read_text())
                while not saved['results'] and time.monotonic() < deadline:
                    time.sleep(0.01)
                    saved = json.loads((results / 'results.json').read_text())
                self.assertEqual(len(saved['results']), 1)
                raise RuntimeError('interrupted')
            return original(task, *args, **kwargs)

        arguments = ['--tasks', ','.join(task['id'] for task in tasks), self.command(), results, 'test-model']
        with patch.dict(os.environ, {'FAKE_MODE': 'no-tools'}), patch.object(runner, 'run_task', side_effect=crash):
            with self.assertRaisesRegex(RuntimeError, 'interrupted'):
                self.invoke(*arguments)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual(len(report['results']), 1)
        self.assertFalse(report['results'][0]['pass'])
        with patch.object(runner, 'run_task', wraps=original) as run:
            self.assertEqual(self.invoke('--resume', *arguments), 1)
        self.assertEqual(run.call_count, 1)
        self.assertEqual(run.call_args.args[0]['id'], tasks[1]['id'])

    def test_v3_cross_model_and_task_filters(self):
        tasks = [dict(task, suite='v3', cross_model=index == 0) for index, task in enumerate(runner.checker.tasks()[:3])]
        results = self.root / 'cross-model'
        fixture, check = runner.checker.fixture, runner.checker.check
        with patch.object(runner.checker, 'tasks', return_value=tasks) as available, \
                patch.object(runner.checker, 'fixture', side_effect=lambda task: fixture(dict(task, suite='v2'))), \
                patch.object(runner.checker, 'check', side_effect=lambda task, output: check(dict(task, suite='v2'), output)):
            self.assertEqual(self.invoke('--suite', 'v3', '--cross-model', '--tasks', ','.join(task['id'] for task in tasks[:2]),
                                         self.command(), results, 'test-model'), 0)
        available.assert_called_once_with(suite='v3')
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual(report['suite'], 'v3')
        self.assertEqual([record['task'] for record in report['results']], [tasks[0]['id']])
        self.assertTrue((results / tasks[0]['id'] / 'run-1' / tasks[0]['output']).is_file())
        for arguments in [('--tasks', 'unknown'), ('--runs', '0'), ('--jobs', '5'), ('--cross-model',)]:
            with self.subTest(arguments=arguments), self.assertRaises(SystemExit) as error:
                self.invoke(*arguments, self.command(), self.root / 'invalid')
            self.assertEqual(error.exception.code, 2)

    def test_v3_tier_selection_and_report(self):
        basic = dict(runner.checker.tasks()[0], suite='v3')
        complex_task = dict(basic, id='xlsx-complex-runner', tier='complex')
        def available(suite='v2', tier='basic'):
            self.assertEqual(suite, 'v3')
            return [basic] if tier == 'basic' else [complex_task] if tier == 'complex' else [basic, complex_task]
        def run(task, cli, model, results, arm, python, index, **kwargs):
            return {'task': task['id'], 'format': task['format'], 'run': index,
                    'pass': True, 'passed': True, 'check_message': 'verified', 'tool_calls': 1,
                    'duration_seconds': 0.01, 'seconds': 0.01, 'agent_exit_code': 0, 'timed_out': False}
        for tier, ids in [('basic', [basic['id']]), ('complex', [complex_task['id']]), ('all', [basic['id'], complex_task['id']])]:
            results = self.root / ('tier-' + tier)
            with self.subTest(tier=tier), patch.object(runner.checker, 'tasks', side_effect=available), \
                    patch.object(runner, 'mcp_ready', return_value=0), patch.object(runner, 'run_task', side_effect=run):
                arguments = [] if tier == 'basic' else ['--tier', tier]
                self.assertEqual(self.invoke('--suite', 'v3', *arguments, self.command(), results, 'test-model'), 0)
                report = json.loads((results / 'results.json').read_text())
                self.assertEqual(report['tier'], tier)
                self.assertEqual(report['task_ids'], ids)
                self.assertEqual([record['task'] for record in report['results']], ids)
        with self.assertRaises(ValueError):
            runner.suite_tasks('v2', 'complex')

    def test_answer_artifact_without_document_output(self):
        from types import SimpleNamespace
        base = runner.checker.tasks()[0]
        source = runner.checker.fixture(base)
        task = dict(base, id='xlsx-complex-answer-runner', suite='v3', tier='complex',
                    expected={'answer': {'width': 10, 'hidden': False}})
        results = self.root / 'answer-results'
        results.mkdir()
        def process(command, cwd, stdout, **kwargs):
            (Path(cwd) / 'answer.json').write_text(json.dumps(task['expected']['answer']))
            stdout.write(json.dumps({'type': 'item.completed', 'item': {'id': 'answer-write',
                         'type': 'file_change', 'changes': [{'path': 'answer.json'}]}}) + '\n')
            return SimpleNamespace(returncode=0, stdin=None, stdout=None, stderr=None,
                                   poll=lambda: 0, wait=lambda timeout: 0)
        with patch.object(runner.checker, 'fixture', return_value=source), \
                patch.object(runner, 'command_line', return_value=['stand-in']), \
                patch.object(runner.subprocess, 'Popen', side_effect=process):
            record = runner.run_task(task, None, 'test-model', results, arm='generic', python=self.python)
        self.assertTrue(record['pass'], record['check_message'])
        self.assertEqual(record['tool_calls'], 1)
        output = results / task['id']
        self.assertEqual(json.loads((output / 'answer.json').read_text()), task['expected']['answer'])
        self.assertFalse((output / task['output']).exists())

    def test_answer_from_final_message(self):
        from types import SimpleNamespace
        base = runner.checker.tasks()[0]
        source = runner.checker.fixture(base)
        task = dict(base, id='xlsx-complex-answer-reply', suite='v3', tier='complex', expected={'answer': 4})
        results = self.root / 'answer-reply-results'
        results.mkdir()
        for reply, passed in [('```json\n4\n```', True), ('4', True), ('The heading starts on page 4.', False)]:
            def process(command, cwd, stdout, **kwargs):
                for item in [{'id': 'call', 'type': 'mcp_tool_call', 'server': 'betteroffice'},
                             {'id': 'early', 'type': 'agent_message', 'text': '5'},
                             {'id': 'final', 'type': 'agent_message', 'text': reply}]:
                    stdout.write(json.dumps({'type': 'item.completed', 'item': item}) + '\n')
                return SimpleNamespace(returncode=0, stdin=None, stdout=None, stderr=None,
                                       poll=lambda: 0, wait=lambda timeout: 0)
            with patch.object(runner.checker, 'fixture', return_value=source), \
                    patch.object(runner, 'command_line', return_value=['stand-in']), \
                    patch.object(runner.subprocess, 'Popen', side_effect=process):
                record = runner.run_task(task, ['stand-in'], 'test-model', results)
            self.assertEqual(record['pass'], passed, record['check_message'])
            self.assertTrue((results / task['id'] / 'answer.json').is_file())

    def test_fixture_sha256_is_reverified_before_each_run(self):
        results = self.root / 'pins'
        results.mkdir()
        task = runner.checker.tasks()[0]
        fixture = self.root / 'pinned.xlsx'
        fixture.write_bytes(runner.checker.fixture(task).read_bytes())
        digest = hashlib.sha256(fixture.read_bytes()).hexdigest()

        def verified(task):
            if hashlib.sha256(fixture.read_bytes()).hexdigest() != digest:
                raise ValueError('sha256 mismatch')
            return fixture

        with patch.object(runner.checker, 'fixture', side_effect=verified) as verify, \
                patch.object(runner.checker, 'check'):
            first = runner.run_task(task, self.cli, 'test-model', results, run=1)
            fixture.write_bytes(b'tampered')
            second = runner.run_task(task, self.cli, 'test-model', results, run=2)
        self.assertTrue(first['pass'], first['reason'])
        self.assertFalse(second['pass'])
        self.assertEqual(second['reason'], 'fixture verification failed: sha256 mismatch')
        self.assertEqual(second['tool_calls'], 0)
        self.assertEqual(verify.call_count, 2)
        self.assertFalse((results / task['id'] / 'run-2' / 'codex.jsonl').exists())

    def test_scripted_bo_replays_examples_with_fresh_servers_and_bound_handles(self):
        server = executable(self.bin / 'scripted-mcp', SCRIPT_MCP)
        results, calls = self.root / 'scripted-bo', self.root / 'scripted-calls.jsonl'
        ids = 'docx-01-fix-heading,xlsx-04-style-header'
        with patch.dict(os.environ, {'SCRIPT_CALLS': str(calls)}), \
                patch.object(runner.shutil, 'which', side_effect=lambda name: None if name == 'codex' else str(server)):
            self.assertEqual(self.invoke('--arm', 'scripted-bo', '--tasks', ids, '--runs', '2', '--jobs', '2', str(server), results), 0)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual((report['agent'], report['model']), ('none', 'none'))
        self.assertEqual([record['tool_calls'] for record in report['results']], [3, 3, 4, 4])
        roots = {call['root'] for call in map(json.loads, calls.read_text().splitlines())}
        self.assertEqual(len(roots), 4)
        self.assertTrue(all(record['pass'] for record in report['results']))
        for record in report['results']:
            self.assertTrue((results / record['task'] / f'run-{record["run"]}' / 'script.jsonl').is_file())

    def test_scripted_bo_stops_on_tool_error(self):
        server = executable(self.bin / 'scripted-mcp', SCRIPT_MCP)
        results, calls = self.root / 'script-error', self.root / 'script-error-calls.jsonl'
        results.mkdir()
        task = next(task for task in runner.checker.tasks() if task['id'] == 'docx-01-fix-heading')
        with patch.dict(os.environ, {'SCRIPT_CALLS': str(calls), 'SCRIPT_FAIL': 'office_grep'}):
            record = runner.run_task(task, [str(server)], 'none', results, arm='scripted-bo')
        self.assertFalse(record['pass'])
        self.assertIn('office_grep failed', record['reason'])
        self.assertEqual(record['tool_calls'], 2)
        self.assertEqual([json.loads(line)['name'] for line in calls.read_text().splitlines()], ['office_open', 'office_grep'])

    def test_scripted_lib_harness_uses_selected_python_and_task_cwd(self):
        python = executable(self.root / 'library-python', '''#!/usr/bin/env python3
import json
import os
import shutil
import sys
from pathlib import Path
if sys.argv[1] == '-c':
    sys.exit(0)
script = Path(sys.argv[1])
assert script.is_file() and script.parent.name == 'lib'
assert Path.cwd().name.startswith('office-format-')
assert len(list(Path.cwd().iterdir())) == 1
ext = '.docx' if script.stem.startswith('docx') else '.xlsx'
shutil.copyfile(Path(os.environ['FAKE_CORRECT']) / (script.stem + ext), 'output' + ext)
print(json.dumps({'type': 'script.call', 'name': 'library.edit'}))
''')
        results = self.root / 'scripted-lib'
        with patch.object(runner.shutil, 'which', return_value=None):
            self.assertEqual(self.invoke('--arm', 'scripted-lib', '--python', python,
                                         '--tasks', 'docx-01-fix-heading,xlsx-04-style-header', results), 0)
        report = json.loads((results / 'results.json').read_text())
        self.assertEqual((report['agent'], report['model']), ('none', 'none'))
        self.assertTrue(all(record['passed'] and record['tool_calls'] == 1 for record in report['results']))

    def test_scripted_library_examples(self):
        try:
            runner.python_ready(Path(sys.executable))
        except RuntimeError as error:
            self.skipTest(str(error))
        results = self.root / 'library-examples'
        results.mkdir()
        for task in runner.checker.tasks():
            if task['id'] in ('docx-01-fix-heading', 'xlsx-04-style-header'):
                with self.subTest(task=task['id']):
                    with patch.object(runner.checker, 'check', wraps=runner.checker.check) as check:
                        record = runner.run_task(task, None, 'none', results, arm='scripted-lib', python=Path(sys.executable))
                    check.assert_called_once()
                    self.assertEqual(record['agent_exit_code'], 0)
                    self.assertEqual(record['tool_calls'], 3 if task['format'] == 'docx' else 5)
                    output = results / task['id'] / task['output']
                    if task['format'] == 'docx':
                        from docx import Document
                        texts = [paragraph.text for paragraph in Document(output).paragraphs]
                        self.assertIn('Summary', texts)
                        self.assertNotIn('Sumary', texts)
                    else:
                        from openpyxl import load_workbook
                        header = load_workbook(output)['Sales'][1][:3]
                        self.assertTrue(all(cell.font.bold and cell.fill.fgColor.rgb == 'FFD9EAF7' for cell in header))

    def test_missing_script_records_failure(self):
        for arm in ('scripted-bo', 'scripted-lib'):
            results = self.root / arm
            results.mkdir()
            with patch.object(runner, 'SCRIPTED', self.root / 'no-scripts'):
                record = runner.run_task(runner.checker.tasks()[0], self.cli, 'none', results, arm=arm, python=self.python)
            self.assertFalse(record['passed'])
            self.assertEqual(record['reason'], 'no script')
            self.assertEqual(record['tool_calls'], 0)

    def test_muse_is_refused_without_starting_a_run(self):
        with patch.object(runner.subprocess, 'Popen', side_effect=AssertionError('must not start Muse')), \
                self.assertRaises(SystemExit) as error:
            self.invoke('--agent', 'muse', self.command(), self.root / 'muse')
        self.assertEqual(error.exception.code, 2)

    def test_atomic_results_survive_replace_failure(self):
        path = self.root / 'atomic.json'
        path.write_text('{"old": true}')
        with patch.object(runner.os, 'replace', side_effect=OSError('disk error')):
            with self.assertRaisesRegex(OSError, 'disk error'):
                runner.atomic_write(path, '{"new": true}')
        self.assertEqual(json.loads(path.read_text()), {'old': True})
        self.assertEqual(list(self.root.glob('.atomic.json*')), [])


class InputPreparationTests(unittest.TestCase):
    def test_unprepared_input_is_byte_identical(self):
        data = b'unchanged fixture bytes'
        self.assertIs(runner.checker.prepare_input({}, data), data)
        self.assertIs(runner.checker.prepare_input({'prepare': []}, data), data)

    def test_unknown_preparation_fails(self):
        with self.assertRaisesRegex(ValueError, 'unknown prepare step: unknown'):
            runner.checker.prepare_input({'prepare': ['unknown']}, b'')

    def test_remove_sheet_protection_preserves_other_bytes(self):
        protection = b'<sheetProtection sheet="1" objects="1" scenarios="1"/>'
        sheet = b'<?xml version="1.0"?><worksheet><sheetData/>\n' + protection + b'<pageMargins/></worksheet>'
        parts = {'xl/worksheets/sheet1.xml': sheet,
                 'xl/worksheets/sheet2.xml': sheet.replace(protection, b'<sheetProtection\n sheet="1" />'),
                 'xl/worksheets/sheet3.xml': b'<worksheet><sheetData/></worksheet>',
                 'xl/worksheets/_rels/sheet1.xml.rels': protection,
                 'xl/worksheets/nested/other.xml': protection,
                 'xl/workbook.xml': protection, 'media/image.bin': bytes(range(256))}
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as archive:
            archive.comment = b'archive metadata'
            for index, (name, content) in enumerate(parts.items()):
                entry = ZipInfo(name, date_time=(2001, 2, 3, 4, 5, 6))
                entry.compress_type = ZIP_DEFLATED if index % 2 else ZIP_STORED
                entry.comment = b'entry metadata'
                entry.external_attr = 0o100644 << 16
                archive.writestr(entry, content)
        data = buffer.getvalue()
        task = {'prepare': ['remove_sheet_protection']}
        prepared = runner.checker.prepare_input(task, data)
        with ZipFile(BytesIO(data)) as source, ZipFile(BytesIO(prepared)) as actual:
            self.assertEqual(source.namelist(), actual.namelist())
            self.assertEqual(source.comment, actual.comment)
            for name, content in parts.items():
                expected = content
                if name in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
                    expected = b'<?xml version="1.0"?><worksheet><sheetData/>\n<pageMargins/></worksheet>'
                self.assertEqual(actual.read(name), expected)
                for attr in ('date_time', 'compress_type', 'comment', 'external_attr'):
                    self.assertEqual(getattr(source.getinfo(name), attr), getattr(actual.getinfo(name), attr))
        self.assertEqual(runner.checker.prepare_input(task, prepared), prepared)

    def test_only_protected_sheet_tasks_prepare_input(self):
        prepared = [task for task in runner.checker.tasks('v3', 'all') if task.get('prepare')]
        self.assertEqual([task['id'] for task in prepared], ['xlsx-v3-cx-07-green-recovery-sort', 'xlsx-v3-cx-08-funding-share', 'xlsx-v3-cx-09-balance-flag'])
        for task in prepared:
            self.assertEqual(task['prepare'], ['remove_sheet_protection'])
            self.assertTrue(task['note'])


class CommentContentTypeTests(unittest.TestCase):
    def package(self, kind, content_type, default=False):
        node = 'Default Extension="xml"' if default else f'Override PartName="/word/{kind}.xml"'
        types = (f'<Types xmlns="{runner.checker.CT_NS}"><{node} ContentType="{content_type}"/></Types>')
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as archive:
            archive.writestr('[Content_Types].xml', types)
            archive.writestr(f'word/{kind}.xml', '<root/>')
        return ZipFile(BytesIO(buffer.getvalue()))

    def test_comment_content_types_are_equivalent(self):
        checker = runner.checker
        for kind in ('comments', 'commentsExtended', 'commentsIds', 'commentsExtensible'):
            for default in (False, True):
                with self.subTest(kind=kind, default=default):
                    modern = f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml'
                    legacy = f'application/vnd.ms-word.{kind}+xml'
                    with self.package(kind, modern, default) as expected, self.package(kind, legacy, default) as actual:
                        self.assertEqual(checker.content_type(actual, f'word/{kind}.xml'), modern)
                        self.assertEqual(checker.content_type(expected, f'word/{kind}.xml'), modern)
                        self.assertEqual(checker.package_state(expected, set(), {}), checker.package_state(actual, set(), {}))
                        self.assertEqual(checker.real_package_state(checker.RealPackage(expected)),
                                         checker.real_package_state(checker.RealPackage(actual)))

    def test_other_content_types_still_differ(self):
        checker = runner.checker
        for modern, legacy in (
                ('application/vnd.openxmlformats-officedocument.wordprocessingml.document+xml',
                 'application/vnd.ms-word.document+xml'),
                ('application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml',
                 'application/vnd.ms-word.commentsExtended+xml')):
            with self.subTest(modern=modern, legacy=legacy):
                with self.package('comments', modern) as expected, self.package('comments', legacy) as actual:
                    self.assertNotEqual(checker.content_type(expected, 'word/comments.xml'),
                                        checker.content_type(actual, 'word/comments.xml'))
                    self.assertNotEqual(checker.real_package_state(checker.RealPackage(expected)),
                                        checker.real_package_state(checker.RealPackage(actual)))

    def test_synthesized_extensions_use_openxmlformats_type(self):
        checker = runner.checker
        modern = 'application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml'
        with self.package('comments', 'application/xml') as source:
            package = checker.RealPackage(source)
        with patch.object(checker, 'related', side_effect=['word/document.xml', None]):
            checker.comment_extensions(package, 'word/document.xml')
        entry = next(node for node in package.root('[Content_Types].xml')
                     if node.get('PartName') == '/word/commentsExtended1.xml')
        self.assertEqual(entry.get('ContentType'), modern)


class AggregateTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='office-aggregate-test-')
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def report(self, name, arm, outcomes, **metadata):
        records = []
        for task, values in outcomes.items():
            for index, passed in enumerate(values, 1):
                number = len(records) + 1
                records.append({'task': task, 'run': index, 'pass': passed,
                                'tool_calls': number, 'duration_seconds': 2 * number})
        path = self.root / name
        path.write_text(json.dumps({'suite': 'v3', 'arm': arm, 'model': 'test-model', 'agent': 'codex',
                                    'runs': 3, **metadata, 'results': records}))
        return path

    def test_hand_computed_statistics_and_paired_intervals(self):
        betteroffice = self.report('bo.json', 'betteroffice', {'a': [True] * 3, 'b': [True, False, False], 'c': [False] * 3})
        generic = self.report('lib.json', 'generic', {task: [True, False, False] for task in ('a', 'b', 'c')})
        summary = aggregator.aggregate([betteroffice, generic])
        bo, lib = summary['groups']
        self.assertEqual((bo['tasks'], bo['runs'], bo['task_runs'], bo['pass_count']), (3, 3, 9, 4))
        self.assertAlmostEqual(bo['pass_rate'], 4 / 9)
        self.assertEqual(bo['bootstrap_95'], [0, 1])
        self.assertAlmostEqual(bo['wilson_95'][0], 0.1888, delta=0.0001)
        self.assertAlmostEqual(bo['wilson_95'][1], 0.7333, delta=0.0001)
        self.assertEqual(bo['all_runs_pass_rate'], 1 / 3)
        self.assertEqual(bo['any_pass_rate'], 2 / 3)
        self.assertEqual([task['pass_count'] for task in bo['per_task']], [3, 1, 0])
        self.assertAlmostEqual(bo['per_task'][1]['variance'], 2 / 9)
        self.assertEqual((bo['median_tool_calls'], bo['median_seconds']), (5, 10))
        self.assertEqual(lib['bootstrap_95'], [1 / 3, 1 / 3])
        comparison, = summary['comparisons']
        self.assertEqual((comparison['arm_a'], comparison['arm_b']), ('betteroffice', 'generic'))
        self.assertAlmostEqual(comparison['difference'], 1 / 9)
        self.assertAlmostEqual(comparison['bootstrap_95'][0], -1 / 3)
        self.assertAlmostEqual(comparison['bootstrap_95'][1], 2 / 3)
        self.assertEqual(summary['bootstrap_seed'], 0)
        self.assertEqual(summary['bootstrap_resamples'], 10000)
        self.assertEqual(summary, aggregator.aggregate([generic, betteroffice]))

    def test_aggregate_cli_is_deterministic_and_writes_markdown(self):
        path = self.report('runs.json', 'generic', {'a': [True, False, True], 'b': [False] * 3})
        first, second, md = self.root / 'first.json', self.root / 'second.json', self.root / 'table.md'
        for out in (first, second):
            result = subprocess.run([sys.executable, str(ROOT / 'aggregate.py'), str(path), '--out', str(out), '--md', str(md)],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(first.read_bytes(), second.read_bytes())
        self.assertIn('| v3 | generic | test-model | codex | 2 | 3 |', md.read_text())
        self.assertIn('Cluster 95%', md.read_text())

    def test_wilson_bounds_for_zero_and_all_passes(self):
        for successes in (0, 1, 9):
            lower, upper = aggregator.wilson(successes, 9)
            self.assertLessEqual(0, lower)
            self.assertLess(lower, upper)
            self.assertLessEqual(upper, 1)
            self.assertLessEqual(lower, successes / 9)
            self.assertGreaterEqual(upper, successes / 9)
        self.assertAlmostEqual(aggregator.wilson(0, 9)[1], 0.2991, delta=0.0001)
        self.assertAlmostEqual(aggregator.wilson(9, 9)[0], 0.7009, delta=0.0001)

    def test_legacy_records_shards_and_mismatched_task_sets(self):
        path = self.report('old.json', 'betteroffice', {'a': [True]})
        old = json.loads(path.read_text())
        for field in ('suite', 'runs', 'agent'):
            old.pop(field)
        old['results'][0].pop('run')
        path.write_text(json.dumps(old))
        shard = self.report('shard.json', 'betteroffice', {'b': [False]}, suite='v2')
        generic = self.report('generic.json', 'generic', {'a': [True]}, suite='v2')
        summary = aggregator.aggregate([path, shard, generic])
        self.assertEqual(summary['groups'][0]['task_runs'], 2)
        self.assertEqual(summary['comparisons'], [])
        with self.assertRaisesRegex(ValueError, 'duplicate task/run'):
            aggregator.aggregate([path, path])

    def test_group_keys_include_suite_model_and_agent(self):
        paths = [self.report(f'{index}.json', 'generic', {'a': [True]}, **metadata) for index, metadata in enumerate([
            {}, {'suite': 'v2'}, {'model': 'another-model'}, {'agent': 'another-agent'}])]
        self.assertEqual(len(aggregator.aggregate(paths)['groups']), 4)


class FixtureManifestTests(unittest.TestCase):
    def test_provenance_sits_beside_the_pins(self):
        manifest = json.loads((ROOT / 'fixtures.json').read_text())
        self.assertEqual(set(manifest), {'origin', 'provenance', 'fixtures'})
        self.assertIn('eval/formats/generate-fixtures.ts', manifest['provenance'])
        tasks = runner.checker.tasks()
        self.assertEqual(set(manifest['fixtures']), {task['fixture'] for task in tasks})
        self.assertTrue(all(set(entry) == {'bytes', 'sha256'} for entry in manifest['fixtures'].values()))
        self.assertTrue(all(runner.checker.fixture(task).is_file() for task in tasks))


if __name__ == '__main__':
    unittest.main()
